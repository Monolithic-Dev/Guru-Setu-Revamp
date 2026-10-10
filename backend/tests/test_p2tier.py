"""P2-tier features: session feedback (API-042), learner level and at-risk mark (API-043), suggested at-risk from
academics (REQ-029), faculty track record (API-044) and MET-11 in the mentoring pack. Uses the test_tracking graph."""
from io import BytesIO

from openpyxl import load_workbook

from tests.test_tracking import (
    NOTE,
    cypher,
    session,
    world,  # noqa: F401  (autouse fixture: the base graph)
)


def _patch_mentee(client, headers, sid="u-aarav", **body):
    return client.patch(f"/api/v1/mentees/{sid}", headers=headers, json=body)


# ---- API-042 session feedback -----------------------------------------------------------------------------------------


def test_feedback_replaces_and_is_author_only(client, aarav, diya, meera):
    session("s-fb", "u-meera", "u-aarav", "2026-10-01", notes=NOTE)
    session("s-abs", "u-meera", "u-aarav", "2026-10-02", present=False)
    url = "/api/v1/sessions/s-fb/feedback"
    assert client.post(url, headers=aarav, json={"feedback": "helpful", "comment": " clear "}).status_code == 200
    first = client.get("/api/v1/sessions/s-fb", headers=meera).json()["attendees"][0]
    assert (first["feedback"], first["feedback_comment"]) == ("helpful", "clear")
    client.post(url, headers=aarav, json={"feedback": "not_helpful"})  # the second submission replaces the first
    second = client.get("/api/v1/sessions/s-fb", headers=meera).json()["attendees"][0]
    assert (second["feedback"], second["feedback_comment"]) == ("not_helpful", None)
    mine = client.get("/api/v1/sessions/s-fb", headers=aarav).json()["attendees"][0]
    assert "feedback" not in mine and "feedback_comment" not in mine  # the student never reads it back
    absent = client.post("/api/v1/sessions/s-abs/feedback", headers=aarav, json={"feedback": "helpful"})
    assert absent.status_code == 409 and absent.json()["code"] == "NOT_PRESENT"
    assert client.post(url, headers=diya, json={"feedback": "helpful"}).status_code == 404  # not an attendee
    assert client.post(url, headers=meera, json={"feedback": "helpful"}).status_code == 403
    assert client.post(url, headers=aarav, json={"feedback": "great"}).status_code == 422
    assert client.post(url, headers=aarav, json={"feedback": "helpful", "comment": "x" * 301}).status_code == 422


# ---- API-043 learner level, note, at-risk mark ------------------------------------------------------------------------


def test_learner_level_note_and_at_risk(client, meera, arjun, suresh, aarav):
    r = _patch_mentee(client, meera, learner_level="slow", learner_note="Needs weekly check-ins", marked_at_risk=True)
    assert r.status_code == 200
    body = r.json()
    assert (body["learner_level"], body["learner_note"], body["marked_at_risk"]) == (
        "slow", "Needs weekly check-ins", True,
    )
    assert "marked_at_risk" in body["attention_reasons"]
    home = {x["student_id"]: x for x in client.get("/api/v1/me/mentees", headers=meera).json()}
    assert "marked_at_risk" in home["u-aarav"]["attention_reasons"]
    # the HOD sees the level only, never the note or the at-risk mark
    hod = client.get("/api/v1/mentees/u-aarav/overview", headers=suresh).json()
    assert hod["learner_level"] == "slow" and "learner_note" not in hod and "marked_at_risk" not in hod
    assert "Needs weekly" not in client.get("/api/v1/me/mentorship", headers=aarav).text  # nor does the student
    # clearing works, and a partial update leaves the other fields alone
    cleared = _patch_mentee(client, meera, learner_level=None).json()
    assert cleared["learner_level"] is None and cleared["learner_note"] == "Needs weekly check-ins"
    assert _patch_mentee(client, meera, marked_at_risk=False).json()["marked_at_risk"] is False
    # only the ASSIGNED mentor may set them: another mentor and even a guide get 404
    assert _patch_mentee(client, arjun, learner_level="slow").status_code == 404
    cypher("MATCH (f:Faculty {user_id: 'u-arjun'}), (s:Student {user_id: 'u-aarav'}) "
           "CREATE (f)-[:MENTORS {id: 'g-p2', type: 'guide', status: 'active', since: '2026-06-01'}]->(s)")
    assert _patch_mentee(client, arjun, learner_level="slow").status_code == 404
    assert _patch_mentee(client, suresh, learner_level="slow").status_code == 403
    assert _patch_mentee(client, aarav, learner_level="slow").status_code == 403
    assert _patch_mentee(client, meera).status_code == 422
    assert _patch_mentee(client, meera, learner_level="genius").status_code == 422
    assert _patch_mentee(client, meera, learner_note="x" * 501).status_code == 422


# ---- REQ-029 suggested at-risk -----------------------------------------------------------------------------------------


def test_suggested_at_risk_from_academics(client, meera):
    cypher("CREATE (:Term {id: '2026-EVEN', name: 'Even 2026', start_date: '2026-01-05', end_date: '2026-05-31'})")
    cypher("MATCH (:Student)-[:HAS_RECORD]->(old:AcademicRecord) DETACH DELETE old")
    cypher(
        """
        MATCH (a:Student {user_id: 'u-aarav'}), (o:Student {user_id: 'u-t-old'}), (k:Student {user_id: 'u-t-ok'})
        CREATE (a)-[:HAS_RECORD]->(:AcademicRecord {id: 'a1', term_id: '2026-ODD', cgpa: 8.0, attendance_pct: 68.0,
                backlogs: 0})
        CREATE (o)-[:HAS_RECORD]->(:AcademicRecord {id: 'o0', term_id: '2026-EVEN', cgpa: 8.6, attendance_pct: 90.0,
                backlogs: 0})
        CREATE (o)-[:HAS_RECORD]->(:AcademicRecord {id: 'o1', term_id: '2026-ODD', cgpa: 7.5, attendance_pct: 90.0,
                backlogs: 0})
        CREATE (k)-[:HAS_RECORD]->(:AcademicRecord {id: 'k0', term_id: '2026-EVEN', cgpa: 8.0, attendance_pct: 90.0,
                backlogs: 0})
        CREATE (k)-[:HAS_RECORD]->(:AcademicRecord {id: 'k1', term_id: '2026-ODD', cgpa: 7.1, attendance_pct: 80.0,
                backlogs: 0})
        """
    )
    rows = {r["student_id"]: r for r in client.get("/api/v1/me/mentees", headers=meera).json()}
    assert "suggested_at_risk" in rows["u-aarav"]["attention_reasons"]
    assert rows["u-aarav"]["attention_details"] == {"suggested_at_risk": "attendance 68%"}
    assert rows["u-t-old"]["attention_details"] == {"suggested_at_risk": "CGPA 8.6 to 7.5"}  # a drop of 1.1
    assert "suggested_at_risk" not in rows["u-t-ok"]["attention_reasons"]  # a 0.9 drop with 80% attendance is fine


# ---- API-044 faculty track record -------------------------------------------------------------------------------------


def test_track_record_is_facts_only(client, aarav, meera, suresh):
    cypher(
        """
        MATCH (f:Faculty {user_id: 'u-meera'}), (a:Student {user_id: 'u-aarav'}), (t:Student {user_id: 'u-t-ok'})
        CREATE (o1:Opening {id: 'o-1', title: 'One', status: 'open', created_at: '2026-08-01T00:00:00.000Z'})
        CREATE (o2:Opening {id: 'o-2', title: 'Two', status: 'open', created_at: '2026-08-01T00:00:00.000Z'})
        CREATE (f)-[:POSTED]->(o1) CREATE (f)-[:POSTED]->(o2)
        CREATE (a)-[:APPLIED_TO {id: 'ap-1', status: 'accepted', applied_at: '2026-09-01T00:00:00.000Z'}]->(o1)
        CREATE (t)-[:APPLIED_TO {id: 'ap-2', status: 'rejected', applied_at: '2026-09-01T00:00:00.000Z'}]->(o2)
        CREATE (f)-[:MENTORS {id: 'g1', type: 'guide', status: 'ended', end_reason: 'completed',
                since: '2026-05-01'}]->(a)
        CREATE (f)-[:MENTORS {id: 'g2', type: 'guide', status: 'active', since: '2026-09-10'}]->(t)
        CREATE (:StatusEvent {id: 'e1', entity_type: 'application', entity_id: 'ap-1', from_status: 'pending',
                to_status: 'accepted', at: '2026-09-03T00:00:00.000Z'})
        CREATE (:StatusEvent {id: 'e2', entity_type: 'application', entity_id: 'ap-2', from_status: 'pending',
                to_status: 'rejected', at: '2026-09-06T00:00:00.000Z'})
        """
    )
    for headers in (aarav, meera, suresh):
        r = client.get("/api/v1/faculty/u-meera/track-record", headers=headers)
        assert r.status_code == 200
    assert r.json() == {
        "faculty_id": "u-meera", "name": "Meera", "openings_posted": 3, "guide_links_accepted": 2,  # 2 + op-edge-ai
        "guide_links_completed": 1, "median_days_to_first_decision": 3.5,  # median of 2 and 5 days
    }
    assert "rating" not in r.text and "review" not in r.text
    empty = client.get("/api/v1/faculty/u-arjun/track-record", headers=aarav).json()
    assert empty["openings_posted"] == 0 and empty["median_days_to_first_decision"] is None
    assert client.get("/api/v1/faculty/u-aarav/track-record", headers=aarav).status_code == 404  # not faculty
    assert client.get("/api/v1/faculty/nobody/track-record", headers=aarav).status_code == 404
    assert client.get("/api/v1/faculty/u-meera/track-record").status_code == 401


# ---- MET-11 -----------------------------------------------------------------------------------------------------------


def test_met_11_in_the_mentoring_pack(client, meera, suresh):
    for sid in ("u-aarav", "u-t-ok"):
        assert _patch_mentee(client, meera, sid, learner_level="slow").status_code == 200
    assert _patch_mentee(client, meera, "u-t-old", learner_level="advanced").status_code == 200
    wb = load_workbook(BytesIO(client.get("/api/v1/evidence/mentoring.xlsx", headers=suresh).content))
    summary = {row[1]: row[2] for row in wb["Summary"].iter_rows(min_row=2, values_only=True)}
    assert summary["Slow learners identified"] == 2  # AC-033.1
    assert [summary[f"learner_level:{k}"] for k in ("slow", "average", "advanced")] == [2, 0, 1]
