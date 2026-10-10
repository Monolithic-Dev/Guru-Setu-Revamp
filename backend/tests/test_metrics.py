"""Dashboard (TEST-32 on a hand-built graph) and the XLSX exports (TEST-08, TEST-21).

The graph is the one from test_tracking (`world`), plus issues, certificates and an `=HYPERLINK` student name.
Expected numbers are worked out by hand in the comments, not read back from the code under test.
"""
from io import BytesIO

import pytest
from openpyxl import load_workbook

from tests.test_tracking import (
    ISSUE_TEXT,
    NOTE,
    cypher,
    issue,
    world,  # noqa: F401  (autouse fixture: the base graph)
)

EVIL = '=HYPERLINK("http://evil.example","x")'


@pytest.fixture(autouse=True)
def extras(world):  # noqa: F811
    # Omar's academic issue: raised in term, resolved in term, with a resolution text that must never leak.
    issue("i-res", "u-t-old", "u-meera", "2026-09-01T08:00:00.000Z", status="resolved")
    cypher("MATCH (i:Issue {id: 'i-res'}) SET i.category = 'academic', i.resolved_at = '2026-10-02T08:00:00.000Z', "
           "i.resolution = 'RESOLUTION-SECRET'")
    cypher("MATCH (s:Student {user_id: 'u-t-ok'}) SET s.name = $n", n=EVIL)
    cypher(
        """
        MATCH (s:Student {user_id: 'u-t-ok'})
        CREATE (s)-[:SUBMITTED]->(:Evidence {id: 'e-1', status: 'certified', title: 'Smart India Hackathon',
               issuer: 'AICTE', issue_date: '2026-08-10', level: 'national', type: 'hackathon',
               certificate_id: 'C-1', reviewed_by: 'u-meera'})
        CREATE (s)-[:SUBMITTED]->(:Evidence {id: 'e-old', status: 'certified', title: 'Old course', issuer: 'X',
               issue_date: '2025-01-10', level: 'online', type: 'course'})
        CREATE (s)-[:SUBMITTED]->(:Evidence {id: 'e-rev', status: 'needs_review', title: 'Pending', type: 'course'})
        """
    )
    cypher("MATCH (c:Criterion) DETACH DELETE c")
    cypher("CREATE (:Criterion {code: 'NAAC-5.3.1', title: 'Awards and medals', evidence_types: ['hackathon'], "
           "verified: false})")


def test_32_dashboard_hod(client, suresh):
    r = client.get("/api/v1/dept/dashboard", headers=suresh)
    assert r.status_code == 200
    d = r.json()
    assert d["term"] == {"id": "2026-ODD", "name": "Odd 2026"}
    # 5 students (Aarav, Diya, Tara, Omar, Sana), all allotted; 2 mentors with mentees: 5 / 2 = 2.5 -> "1:2"
    assert (d["students_total"], d["students_allotted"], d["unallotted"], d["ratio"]) == (5, 5, 0, "1:2")
    # compliant (>= 2 present meetings): Tara and Sana only -> 2 of 5
    assert d["compliance_pct"] == 40.0
    mentors = {m["name"]: m for m in d["mentors"]}
    assert mentors["Meera"]["mentee_count"] == 4 and mentors["Meera"]["mentees_compliant"] == 2
    assert mentors["Meera"]["compliance_pct"] == 50.0 and mentors["Meera"]["meetings_held"] == 5
    assert mentors["Arjun"]["mentee_count"] == 1 and mentors["Arjun"]["compliance_pct"] == 0.0
    assert mentors["Meera"]["overloaded"] is False
    assert d["overdue_actions"] == 1  # a-old only; a-ok-done is done
    assert d["pending_reviews"] == 1  # e-rev
    # Aarav, Diya (no sessions), Omar (old session + overdue), Sana (stale issue) are flagged; Tara is clean
    assert d["needs_attention_count"] == 4
    flagged = {a["student_id"]: a for a in d["needs_attention_list"]}
    assert set(flagged) == {"u-aarav", "u-diya", "u-t-old", "u-t-stale"}
    assert flagged["u-t-old"]["mentor_name"] == "Meera"
    cats = {c["category"]: c for c in d["issues_by_category"]}
    assert cats["personal"] == {"category": "personal", "raised": 2, "resolved": 0}
    assert cats["academic"] == {"category": "academic", "raised": 1, "resolved": 1}
    assert cats["confidential"]["raised"] == 1 and "category" in cats["confidential"]
    assert d["achievements_by_type"] == [{"type": "hackathon", "count": 1}]  # the 2025 course is outside the year
    assert d["escalated_issues"] == []
    assert ISSUE_TEXT not in r.text and "RESOLUTION-SECRET" not in r.text and NOTE not in r.text


def test_32_dashboard_iqac_has_no_student_names(client, kavya):
    r = client.get("/api/v1/dept/dashboard", headers=kavya)
    assert r.status_code == 200
    d = r.json()
    assert d["needs_attention_count"] == 4
    assert "needs_attention_list" not in d and "escalated_issues" not in d
    for name in ("Aarav", "Diya", "Omar", "Sana", "HYPERLINK", "CB.EN."):
        assert name not in r.text


def test_32_dashboard_roles(client, meera, aarav):
    assert client.get("/api/v1/dept/dashboard", headers=meera).status_code == 403
    assert client.get("/api/v1/dept/dashboard", headers=aarav).status_code == 403
    assert client.get("/api/v1/dept/dashboard").status_code == 401


def test_32_escalated_issue_text_only_when_visible(client, suresh):
    cypher("MATCH (i:Issue {id: 'i-stale'}) SET i.escalated = true")
    cypher("MATCH (i:Issue {id: 'i-conf'}) SET i.escalated = true")  # confidential: must stay hidden
    rows = client.get("/api/v1/dept/dashboard", headers=suresh).json()["escalated_issues"]
    assert [e["issue_id"] for e in rows] == ["i-stale"]
    assert rows[0]["text"] == ISSUE_TEXT and rows[0]["mentor_name"] == "Meera"


def _book(client, headers, path):
    r = client.get(path, headers=headers)
    assert r.status_code == 200
    return r, load_workbook(BytesIO(r.content))


def _all_cells(wb) -> list:
    return [c for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if c is not None]


def test_08_21_mentoring_export(client, suresh, kavya, aarav):
    r, wb = _book(client, suresh, "/api/v1/evidence/mentoring.xlsx")
    assert r.headers["content-disposition"] == 'attachment; filename="gurusetu-mentoring-2026-ODD-DRAFT.xlsx"'
    assert wb.sheetnames == ["README", "Mentors", "Allotment", "Summary", "Mentor-wise"]
    assert wb["README"]["A1"].value == "DRAFT – verify against the official NAAC manual before submission"
    cells = _all_cells(wb)
    assert not any(isinstance(c, str) and c.startswith("=") for c in cells)
    assert "'" + EVIL in cells  # SEC-18: the formula-looking student name is neutralised
    for secret in (NOTE, ISSUE_TEXT, "RESOLUTION-SECRET", "Do the thing"):
        assert secret not in cells
    summary = {row[1]: row[2] for row in wb["Summary"].iter_rows(min_row=2, values_only=True)}
    assert summary["compliance_pct"] == 40.0 and summary["ratio"] == "1:2" and summary["overdue_actions"] == 1
    assert [row[0] for row in wb["Mentor-wise"].iter_rows(min_row=2, values_only=True)] == ["Arjun", "Meera"]
    assert client.get("/api/v1/evidence/mentoring.xlsx", headers=kavya).status_code == 200
    assert client.get("/api/v1/evidence/mentoring.xlsx", headers=aarav).status_code == 403
    assert client.get("/api/v1/evidence/mentoring.xlsx").status_code == 401


def test_08_achievements_export(client, suresh):
    _, wb = _book(client, suresh, "/api/v1/evidence/achievements.xlsx")
    assert wb.sheetnames == ["README", "Achievements", "By type"]
    rows = list(wb["Achievements"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1  # only the in-year certified certificate
    roll, name, title, issuer, issue_date, level, kind, cert_id, status, reviewed_by, code = rows[0]
    assert (name, title, level, kind, status, reviewed_by, code) == (
        "'" + EVIL, "Smart India Hackathon", "national", "hackathon", "certified", "Meera", "NAAC-5.3.1",
    )
    assert list(wb["By type"].iter_rows(min_row=2, values_only=True)) == [("hackathon", "national", 1)]


# ---- API-034 click-to-source ------------------------------------------------------------------------------------------


def _records(client, headers, metric, **params):
    r = client.get(f"/api/v1/dept/metrics/{metric}/records", headers=headers, params=params)
    assert r.status_code == 200, r.text
    return r.json()


def test_a1_totals_equal_the_dashboard(client, suresh):
    d = client.get("/api/v1/dept/dashboard", headers=suresh).json()
    assert _records(client, suresh, "needs_attention")["total"] == d["needs_attention_count"] == 4
    assert _records(client, suresh, "overdue_actions")["total"] == d["overdue_actions"] == 1
    assert _records(client, suresh, "pending_reviews")["total"] == d["pending_reviews"] == 1
    assert _records(client, suresh, "achievements")["total"] == sum(a["count"] for a in d["achievements_by_type"])
    assert _records(client, suresh, "issues_raised")["total"] == sum(c["raised"] for c in d["issues_by_category"])
    assert _records(client, suresh, "issues_resolved")["total"] == sum(c["resolved"] for c in d["issues_by_category"])
    compliance = _records(client, suresh, "compliance")
    assert compliance["total"] == d["students_allotted"]
    assert sum(r["compliant"] for r in compliance["items"]) / compliance["total"] * 100 == d["compliance_pct"]


def test_a1_rows_for_hod_have_no_text(client, suresh):
    raised = _records(client, suresh, "issues_raised")
    assert {r["category"] for r in raised["items"]} == {"personal", "academic", "confidential"}
    row = next(r for r in raised["items"] if r["category"] == "academic")
    assert row["student"] == "Omar Old" and row["roll_no"] and row["student_id"] == "u-t-old" and row["mentor"] == "Meera"
    overdue = _records(client, suresh, "overdue_actions")["items"][0]
    assert overdue["days_overdue"] == 39 and overdue["mentor"] == "Meera"  # 2026-09-01 -> 2026-10-10
    for metric in ("issues_raised", "overdue_actions", "needs_attention", "pending_reviews", "compliance"):
        text = client.get(f"/api/v1/dept/metrics/{metric}/records", headers=suresh).text
        for secret in (ISSUE_TEXT, "RESOLUTION-SECRET", NOTE, "Do the thing"):
            assert secret not in text


def test_a1_iqac_gets_pseudonyms_only(client, kavya):
    for metric in ("compliance", "issues_raised", "overdue_actions", "needs_attention", "pending_reviews"):
        r = client.get(f"/api/v1/dept/metrics/{metric}/records", headers=kavya)
        assert r.status_code == 200
        for forbidden in ("Aarav", "Diya", "Omar", "Sana", "HYPERLINK", "CB.EN.", "u-t-", "u-aarav"):
            assert forbidden not in r.text, (metric, forbidden)
    rows = _records(client, kavya, "needs_attention")["items"]
    assert sorted(r["student"] for r in rows) == ["Student 01", "Student 02", "Student 03", "Student 04"]
    assert all(set(r) == {"student", "mentor", "reasons"} for r in rows)


def test_a1_cursor_pagination_is_stable(client, kavya):
    first = _records(client, kavya, "compliance", limit=2)
    assert len(first["items"]) == 2 and first["next_cursor"] and first["total"] == 5
    second = _records(client, kavya, "compliance", limit=2, cursor=first["next_cursor"])
    third = _records(client, kavya, "compliance", limit=2, cursor=second["next_cursor"])
    assert third["next_cursor"] is None and len(third["items"]) == 1
    labels = [r["student"] for page in (first, second, third) for r in page["items"]]
    assert len(set(labels)) == 5  # no repeats, no gaps, same pseudonyms on every page
    bad = client.get("/api/v1/dept/metrics/compliance/records", headers=kavya, params={"cursor": "***"})
    assert bad.status_code == 422


def test_a1_unknown_metric_and_roles(client, suresh, meera, aarav):
    assert client.get("/api/v1/dept/metrics/nope/records", headers=suresh).status_code == 404
    assert client.get("/api/v1/dept/metrics/compliance/records", headers=meera).status_code == 403
    assert client.get("/api/v1/dept/metrics/compliance/records", headers=aarav).status_code == 403
    assert client.get("/api/v1/dept/metrics/compliance/records").status_code == 401
